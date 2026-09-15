"""Dựng tệp Word / PDF / PowerPoint / Excel từ nội dung bot soạn, để gửi vào nhóm Zalo.

Vì sao cần: người trong nhóm không có công cụ ghi tệp hay chạy lệnh (mở ra là mở cả máy
chủ). Mô-đun này chỉ nhận *nội dung* — chữ Markdown, danh sách slide, bảng — rồi dựng tệp
trong thư mục tạm mà công cụ gọi truyền vào. Không đọc tệp nào trên máy, không tải ảnh,
không chạy lệnh; mọi kích thước đều có trần để một lời nhờ không làm treo bot.
"""

from __future__ import annotations

import os
import re
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

FORMATS = ("docx", "pptx", "xlsx", "pdf")
MAX_TEXT_CHARS = 30_000
MAX_SLIDES = 40
MAX_BULLETS_PER_SLIDE = 15
MAX_SHEETS = 5
MAX_ROWS = 2_000
MAX_COLS = 30
MAX_CELL_CHARS = 1_000
MAX_TITLE_CHARS = 200

# Font có đủ dấu tiếng Việt cho PDF: (thường, đậm). Biến môi trường thắng.
_PDF_FONTS = (
    (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf", "/usr/share/fonts/truetype/noto/NotoSans-Bold.ttf"),
)


class FileSpecError(ValueError):
    """Nội dung không dựng được tệp — thông điệp đưa thẳng cho người nhờ."""


# ---------------------------------------------------------------- nội dung


def safe_filename(name: str, fmt: str) -> str:
    """Tên tệp gọn, giữ chữ tiếng Việt, bỏ ký tự có thể thành đường dẫn."""
    stem = unicodedata.normalize("NFC", str(name or "").strip())
    stem = re.sub(r"\.(docx|pptx|xlsx|pdf)$", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", " ", stem)
    stem = re.sub(r"\s+", "_", stem).strip("._ ")[:80]
    return f"{stem or 'tai_lieu'}.{fmt}"


def parse_blocks(markdown: str) -> List[Tuple]:
    """Tách Markdown đơn giản thành khối: tiêu đề, gạch đầu dòng, đánh số, bảng, đoạn."""
    blocks: List[Tuple] = []
    table: List[List[str]] = []

    def flush_table() -> None:
        if table:
            blocks.append(("table", [row[:] for row in table]))
            table.clear()

    for raw in str(markdown or "").splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|") and len(stripped) > 1:
            cells = [cell.strip() for cell in stripped[1:-1].split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells if cell):
                table.append(cells)
            continue
        flush_table()
        if not stripped:
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        bullet = re.match(r"^([-*•+])\s+(.*)$", stripped)
        number = re.match(r"^(\d+[.)])\s+(.*)$", stripped)
        indent = len(line) - len(line.lstrip())
        if heading:
            blocks.append(("heading", min(len(heading.group(1)), 3), heading.group(2).strip()))
        elif bullet:
            blocks.append(("bullet", 1 if indent >= 2 else 0, bullet.group(2).strip()))
        elif number:
            blocks.append(("number", 0, f"{number.group(1)} {number.group(2).strip()}"))
        else:
            blocks.append(("para", 0, stripped))
    flush_table()
    return blocks


def inline_runs(text: str) -> List[Tuple[str, bool, bool]]:
    """Chia một dòng thành các đoạn (chữ, đậm, nghiêng) theo **đậm** và *nghiêng*."""
    runs: List[Tuple[str, bool, bool]] = []
    for part in re.split(r"(\*\*[^*]+\*\*|(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))", str(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append((part[1:-1], False, True))
        else:
            runs.append((part, False, False))
    return runs


def _plain(text: str) -> str:
    return "".join(chunk for chunk, _b, _i in inline_runs(text))


def _check_text(title: str, content: str) -> None:
    if len(title) > MAX_TITLE_CHARS:
        raise FileSpecError(f"tiêu đề dài quá {MAX_TITLE_CHARS} ký tự")
    if not content.strip():
        raise FileSpecError("cần `content` — nội dung văn bản của tệp")
    if len(content) > MAX_TEXT_CHARS:
        raise FileSpecError(f"nội dung dài quá {MAX_TEXT_CHARS} ký tự — chia thành nhiều tệp nhỏ hơn")


def _normalize_slides(slides: Any) -> List[Dict[str, Any]]:
    if not isinstance(slides, list) or not slides:
        raise FileSpecError("cần `slides` — danh sách slide, mỗi slide có `title` và `bullets`")
    if len(slides) > MAX_SLIDES:
        raise FileSpecError(f"tối đa {MAX_SLIDES} slide mỗi tệp")
    result = []
    for index, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            raise FileSpecError(f"slide {index} phải có dạng {{title, bullets}}")
        bullets = slide.get("bullets") or []
        if not isinstance(bullets, list):
            raise FileSpecError(f"`bullets` của slide {index} phải là danh sách")
        if len(bullets) > MAX_BULLETS_PER_SLIDE:
            raise FileSpecError(f"slide {index} có quá {MAX_BULLETS_PER_SLIDE} ý — tách thành nhiều slide")
        title = str(slide.get("title") or "").strip()[:MAX_TITLE_CHARS]
        items = [str(b).strip()[:MAX_CELL_CHARS] for b in bullets if str(b).strip()]
        if not title and not items:
            raise FileSpecError(f"slide {index} trống")
        result.append({"title": title, "bullets": items})
    return result


def _normalize_sheets(sheets: Any) -> List[Dict[str, Any]]:
    if not isinstance(sheets, list) or not sheets:
        raise FileSpecError("cần `sheets` — danh sách trang tính, mỗi trang có `name` và `rows`")
    if len(sheets) > MAX_SHEETS:
        raise FileSpecError(f"tối đa {MAX_SHEETS} trang tính mỗi tệp")
    result, used = [], set()
    for index, sheet in enumerate(sheets, 1):
        if not isinstance(sheet, dict) or not isinstance(sheet.get("rows"), list) or not sheet["rows"]:
            raise FileSpecError(f"trang tính {index} cần `rows` — danh sách các dòng")
        rows = sheet["rows"]
        if len(rows) > MAX_ROWS:
            raise FileSpecError(f"trang tính {index} quá {MAX_ROWS} dòng")
        clean_rows = []
        for row in rows:
            cells = row if isinstance(row, list) else [row]
            if len(cells) > MAX_COLS:
                raise FileSpecError(f"trang tính {index} quá {MAX_COLS} cột")
            clean_rows.append([_cell(value) for value in cells])
        name = re.sub(r"[\[\]:*?/\\]", " ", str(sheet.get("name") or f"Trang {index}")).strip()[:31] or f"Trang {index}"
        while name in used:
            name = f"{name[:28]} {index}"
        used.add(name)
        result.append({"name": name, "rows": clean_rows})
    return result


def _cell(value: Any) -> Any:
    if isinstance(value, bool) or value is None:
        return "" if value is None else str(value)
    if isinstance(value, (int, float)):
        return value
    text = str(value)[:MAX_CELL_CHARS]
    # Chữ bắt đầu bằng "=" là công thức Excel: người lạ không được cài công thức vào tệp.
    return "'" + text if text.startswith(("=", "+", "-", "@")) and not re.fullmatch(r"[+-]?\d+([.,]\d+)?", text) else text


# ---------------------------------------------------------------- dựng tệp


def build_docx(title: str, content: str, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    if title:
        doc.add_heading(title, 0)
    for block in parse_blocks(content):
        kind = block[0]
        if kind == "table":
            rows = block[1]
            width = max(len(row) for row in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(_plain(row[c]) if c < len(row) else "")
                    run.bold = r == 0
            continue
        if kind == "heading":
            doc.add_heading(_plain(block[2]), block[1])
            continue
        style = {"bullet": "List Bullet 2" if block[1] else "List Bullet", "number": "List Number"}.get(kind)
        text = block[2]
        if kind == "number":
            text = re.sub(r"^\d+[.)]\s+", "", text)
        paragraph = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        for chunk, bold, italic in inline_runs(text):
            run = paragraph.add_run(chunk)
            run.bold, run.italic = bold, italic
            run.font.size = Pt(12)
    doc.save(str(path))


def _pdf_fonts() -> Tuple[str, str]:
    regular = os.environ.get("ZALO_PDF_FONT", "").strip()
    bold = os.environ.get("ZALO_PDF_FONT_BOLD", "").strip() or regular
    candidates = ([(regular, bold)] if regular else []) + list(_PDF_FONTS)
    for reg, bld in candidates:
        if reg and os.path.isfile(reg):
            return reg, bld if bld and os.path.isfile(bld) else reg
    raise FileSpecError("máy chủ thiếu font tiếng Việt để tạo PDF (đặt ZALO_PDF_FONT)")


def build_pdf(title: str, content: str, path: Path) -> None:
    from fpdf import FPDF

    regular, bold = _pdf_fonts()
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("VN", "", regular)
    pdf.add_font("VN", "B", bold)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def write(text: str, size: float, style: str = "", indent: float = 0.0) -> None:
        pdf.set_font("VN", style, size)
        pdf.set_x(pdf.l_margin + indent)
        # fpdf2 hiểu **đậm**; nghiêng không có font riêng nên bỏ dấu *.
        safe = "".join(f"**{chunk}**" if b else chunk for chunk, b, _i in inline_runs(text))
        pdf.multi_cell(width - indent, size * 0.5, safe, markdown=not style, new_x="LMARGIN", new_y="NEXT")

    if title:
        write(title, 18, "B")
        pdf.ln(2)
    for block in parse_blocks(content):
        kind = block[0]
        if kind == "table":
            pdf.set_font("VN", "", 10)
            rows = block[1]
            columns = max(len(row) for row in rows)
            with pdf.table(text_align="LEFT", first_row_as_headings=False) as table:
                for r, row in enumerate(rows):
                    cells = table.row()
                    for c in range(columns):
                        pdf.set_font("VN", "B" if r == 0 else "", 10)
                        cells.cell(_plain(row[c]) if c < len(row) else "")
            pdf.ln(2)
        elif kind == "heading":
            pdf.ln(1)
            write(_plain(block[2]), {1: 15, 2: 13.5, 3: 12.5}[block[1]], "B")
        elif kind == "bullet":
            write("• " + block[2], 11.5, indent=6.0 * (block[1] + 1))
        else:
            write(block[2], 11.5)
    pdf.output(str(path))


def build_pptx(title: str, slides: Sequence[Dict[str, Any]], path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    if title:
        cover = prs.slides.add_slide(prs.slide_layouts[0])
        cover.shapes.title.text = title
        if len(cover.placeholders) > 1:
            cover.placeholders[1].text = ""
    for spec in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = _plain(spec["title"])
        body = slide.placeholders[1].text_frame
        bullets = spec["bullets"]
        size = Pt(24 if len(bullets) <= 5 else 20 if len(bullets) <= 8 else 16)
        body.clear()
        for index, item in enumerate(bullets):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            for chunk, bold, italic in inline_runs(item):
                run = paragraph.add_run()
                run.text = chunk
                run.font.bold, run.font.italic, run.font.size = bold, italic, size
    prs.save(str(path))


def build_xlsx(title: str, sheets: Sequence[Dict[str, Any]], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    book.remove(book.active)
    for spec in sheets:
        sheet = book.create_sheet(spec["name"])
        for row in spec["rows"]:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(8, longest + 2))
    if title:
        book.properties.title = title
    book.save(str(path))


def make_file(
    fmt: str,
    title: str = "",
    *,
    directory: str,
    filename: Optional[str] = None,
    content: Optional[str] = None,
    slides: Any = None,
    sheets: Any = None,
) -> Path:
    """Kiểm tra giới hạn rồi dựng tệp trong ``directory``. Trả đường dẫn tệp đã tạo."""
    fmt = str(fmt or "").lower().lstrip(".")
    if fmt not in FORMATS:
        raise FileSpecError("chỉ tạo được tệp docx, pptx, xlsx hoặc pdf")
    title = str(title or "").strip()
    path = Path(directory) / safe_filename(filename or title, fmt)
    if fmt in ("docx", "pdf"):
        text = str(content or "")
        _check_text(title, text)
        (build_docx if fmt == "docx" else build_pdf)(title, text, path)
    elif fmt == "pptx":
        build_pptx(title[:MAX_TITLE_CHARS], _normalize_slides(slides), path)
    else:
        build_xlsx(title[:MAX_TITLE_CHARS], _normalize_sheets(sheets), path)
    return path
